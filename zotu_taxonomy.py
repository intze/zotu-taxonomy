#!/usr/bin/env python3
"""
zOTU -> species-level taxonomy against 16S rRNA type-strain sequences.

Workflow
  1. Downloads (once) the NCBI RefSeq 16S rRNA "Targeted Loci" collection
     (bacteria + archaea, curated, mostly type material) and vsearch.
  2. Trims leftover degenerate primer bases (auto-detected).
  3. Global alignment of every zOTU with vsearch (--usearch_global).
  4. Keeps the best hit per species, ranks species by % identity and
     reports the top N (default 5) for every zOTU.
  5. Writes an Excel file (one row per zOTU + a long table + settings).

Usage:
  python zotu_taxonomy.py zotus.fasta
  python zotu_taxonomy.py zotus.fasta -o results.xlsx --threads 4
  python zotu_gui.py                       (graphical interface)

Requires Python 3.8+ and openpyxl (pip install openpyxl).
If openpyxl is missing, CSV files are written instead.
"""

import argparse
import csv
import datetime
import gzip
import os
import platform
import re
import shutil
import stat
import subprocess
import sys
import tarfile
import urllib.request
import zipfile
from collections import defaultdict

VSEARCH_VERSION = "2.30.0"
VSEARCH_URL = ("https://github.com/torognes/vsearch/releases/download/"
               "v{v}/vsearch-{v}-{plat}")

DB_URLS = [
    "https://ftp.ncbi.nlm.nih.gov/refseq/TargetedLoci/Bacteria/bacteria.16SrRNA.fna.gz",
    "https://ftp.ncbi.nlm.nih.gov/refseq/TargetedLoci/Archaea/archaea.16SrRNA.fna.gz",
]

# Identity thresholds (Yarza et al. 2014, Nat Rev Microbiol; Kim et al. 2014, IJSEM).
# NOTE: defined for full-length 16S; for short amplicons they are only indicative.
SPECIES_ID = 98.65
GENUS_ID = 94.5

PRIMER_TAILS = [
    # (side, regex, bases to trim, description)
    ("left", r"^GC[ATW]GCAG", 7, "341F tail (CCTACGGGNGGCWGCAG)"),
    ("right", r"ATTAGA[ATW]AC$", 9, "805R/806R tail (GACTACHVGGGTATCTAATCC)"),
]


def app_dir():
    """Folder of the script, or of the .exe when frozen with PyInstaller."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


def data_dir():
    """Where db/ and tools/ live: next to the program if writable, else the user profile."""
    here = app_dir()
    try:
        test = os.path.join(here, ".write_test")
        with open(test, "w") as f:
            f.write("ok")
        os.remove(test)
        return here
    except OSError:
        base = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
        path = os.path.join(base, "ZotuTaxonomy")
        os.makedirs(path, exist_ok=True)
        return path


# --------------------------------------------------------------------------- #
# Downloads
# --------------------------------------------------------------------------- #
def download(url, dest, log=print):
    log(f"  downloading {url}")
    tmp = dest + ".part"
    with urllib.request.urlopen(url, timeout=120) as r, open(tmp, "wb") as f:
        shutil.copyfileobj(r, f)
    os.replace(tmp, dest)


def find_vsearch():
    """vsearch bundled next to the program, in the data folder, or on PATH."""
    exe = "vsearch.exe" if platform.system() == "Windows" else "vsearch"
    for base in {app_dir(), data_dir()}:
        tools = os.path.join(base, "tools")
        if os.path.isdir(tools):
            for root, _dirs, files in os.walk(tools):
                if exe in files:
                    return os.path.join(root, exe)
    return shutil.which("vsearch")


def ensure_vsearch(log=print):
    found = find_vsearch()
    if found:
        return found
    system, machine = platform.system(), platform.machine().lower()
    if system == "Windows":
        plat = "win-x86_64.zip"
    elif system == "Linux":
        plat = "linux-aarch64.tar.gz" if machine in ("aarch64", "arm64") else "linux-x86_64.tar.gz"
    elif system == "Darwin":
        plat = "macos-universal.tar.gz"
    else:
        raise RuntimeError(f"Unsupported OS: {system}. Install vsearch and put it on PATH.")

    tools = os.path.join(data_dir(), "tools")
    os.makedirs(tools, exist_ok=True)
    archive = os.path.join(tools, f"vsearch-{VSEARCH_VERSION}-{plat}")
    log("vsearch not found - downloading it once:")
    download(VSEARCH_URL.format(v=VSEARCH_VERSION, plat=plat), archive, log)
    if archive.endswith(".zip"):
        with zipfile.ZipFile(archive) as z:
            z.extractall(tools)
    else:
        with tarfile.open(archive) as t:
            t.extractall(tools)
    os.remove(archive)
    binary = find_vsearch()
    if not binary:
        raise RuntimeError(f"Could not find vsearch after extracting {archive}")
    os.chmod(binary, os.stat(binary).st_mode | stat.S_IEXEC)
    return binary


def parse_ncbi_name(description):
    """'NR_117678.1 Lactiplantibacillus plantarum strain JCM 1149 16S ...'
    -> 'Lactiplantibacillus plantarum'"""
    words = description.split()[1:]
    if not words:
        return "unknown"
    n = 3 if words[0] == "Candidatus" else 2
    name = []
    for w in words:
        if w in ("strain", "16S", "subsp.", "str.") or len(name) == n:
            break
        name.append(w)
    return " ".join(name).strip("'\"")


def db_paths():
    d = os.path.join(data_dir(), "db")
    return (d, os.path.join(d, "ncbi_16S_refs.fasta"),
            os.path.join(d, "ncbi_16S_names.tsv"), os.path.join(d, "DB_DATE.txt"))


def db_date():
    _d, _f, _n, date_file = db_paths()
    if os.path.exists(date_file):
        with open(date_file) as f:
            return f.read().strip()
    return None


def ensure_db(update=False, log=print):
    """Returns (fasta_path, names_tsv_path). Builds them on first run."""
    db_dir, fasta, names, date_file = db_paths()
    os.makedirs(db_dir, exist_ok=True)
    if os.path.exists(fasta) and os.path.exists(names) and not update:
        return fasta, names

    log("Downloading the NCBI 16S reference database:")
    n = 0
    with open(fasta + ".part", "w") as fo, open(names + ".part", "w", newline="") as fn:
        w = csv.writer(fn, delimiter="\t")
        w.writerow(["accession", "species", "description"])
        for url in DB_URLS:
            gz = os.path.join(db_dir, os.path.basename(url))
            download(url, gz, log)
            with gzip.open(gz, "rt") as fi:
                for line in fi:
                    if line.startswith(">"):
                        desc = line[1:].strip()
                        acc = desc.split()[0]
                        w.writerow([acc, parse_ncbi_name(desc), desc])
                        fo.write(f">{acc}\n")
                        n += 1
                    else:
                        fo.write(line.strip().upper() + "\n")
            os.remove(gz)
    os.replace(fasta + ".part", fasta)
    os.replace(names + ".part", names)
    with open(date_file, "w") as f:
        f.write(datetime.date.today().isoformat() + "\n")
    log(f"  {n} reference sequences")
    return fasta, names


# --------------------------------------------------------------------------- #
# Analysis
# --------------------------------------------------------------------------- #
def read_fasta(path):
    seqs, name = {}, None
    with open(path) as f:
        for line in f:
            line = line.strip()
            if line.startswith(">"):
                name = line[1:].split()[0].split(";")[0]
                seqs[name] = []
            elif name:
                seqs[name].append(line.upper())
    return {k: "".join(v) for k, v in seqs.items()}


def detect_trim(seqs, min_fraction=0.8):
    """Detect leftover degenerate primer bases shared by most zOTUs.
    Degenerate positions (W = A/T) create artificial mismatches against the references."""
    trim = {"left": 0, "right": 0}
    notes = []
    for side, pattern, n, desc in PRIMER_TAILS:
        frac = sum(bool(re.search(pattern, s)) for s in seqs.values()) / max(len(seqs), 1)
        if frac >= min_fraction:
            trim[side] = n
            notes.append(f"{side} {n} bp ({desc}, found in {frac:.0%} of zOTUs)")
    return trim["left"], trim["right"], notes


def run_vsearch(vsearch, query, db, out, threads, min_id, min_cov, maxaccepts, maxrejects,
                log=print, progress=None):
    cmd = [vsearch, "--usearch_global", query, "--db", db,
           "--id", str(min_id / 100), "--query_cov", str(min_cov),
           "--maxaccepts", str(maxaccepts), "--maxrejects", str(maxrejects),
           "--maxhits", str(maxaccepts), "--strand", "both",
           "--userout", out, "--userfields", "query+target+id+alnlen+qcov",
           "--threads", str(threads)]
    log("Running: " + " ".join(cmd))
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)  # no console window on Windows
    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE,
                            stdin=subprocess.DEVNULL, creationflags=flags)
    # vsearch writes progress like "Searching 42%" separated by carriage returns
    buf, last, tail = b"", None, []
    while True:
        chunk = proc.stderr.read(256)
        if not chunk:
            break
        buf += chunk
        parts = re.split(rb"[\r\n]", buf)
        buf = parts.pop()
        for p in parts:
            text = p.decode(errors="replace").strip()
            if not text:
                continue
            tail = (tail + [text])[-20:]
            m = re.match(r"Searching (\d+)%", text)
            if m:
                pct = int(m.group(1))
                if pct != last:
                    last = pct
                    if progress:
                        progress(pct)
            elif not re.search(r"\d+%$", text):
                log("  " + text)
    if proc.wait() != 0:
        raise RuntimeError("vsearch failed:\n" + "\n".join(tail))


def classify(hits_tsv, names_tsv, order, lengths, top_n, species_id, genus_id, tie_delta):
    acc2sp = {}
    with open(names_tsv) as f:
        for row in csv.DictReader(f, delimiter="\t"):
            acc2sp[row["accession"]] = row["species"]

    # best hit per (zOTU, species)
    best = defaultdict(dict)
    with open(hits_tsv) as f:
        for line in f:
            q, t, ident, alnlen, qcov = line.rstrip("\n").split("\t")
            q = q.split(";")[0]
            ident, alnlen, qcov = float(ident), int(alnlen), float(qcov)
            sp = acc2sp.get(t, t)
            cur = best[q].get(sp)
            if cur is None or (ident, alnlen) > (cur[0], cur[1]):
                best[q][sp] = (ident, alnlen, qcov, t)

    results = []
    for q in order:
        ranked = sorted(best.get(q, {}).items(), key=lambda kv: (-kv[1][0], -kv[1][1], kv[0]))
        top = ranked[:top_n]
        row = {"zOTU": q, "length": lengths[q], "hits": top}
        if not ranked:
            row.update(best_id=None, n_tied=0, status="No type strain >= min identity",
                       assigned_species="", assigned_genus="")
            results.append(row)
            continue
        best_id = ranked[0][1][0]
        tied = [sp for sp, v in ranked if best_id - v[0] <= tie_delta]
        genera = sorted({sp.split()[0] for sp in tied})
        genus = genera[0] if len(genera) == 1 else " / ".join(genera)
        if best_id >= species_id:
            if len(tied) == 1:
                status, species = "Species", tied[0]
            else:
                status = f"Ambiguous: {len(tied)} species tied"
                species = " / ".join(tied[:top_n]) + (" / ..." if len(tied) > top_n else "")
        elif best_id >= genus_id:
            status, species = "Genus level only", ""
        else:
            status, species, genus = "Below genus threshold", "", ""
        row.update(best_id=best_id, n_tied=len(tied), status=status,
                   assigned_species=species, assigned_genus=genus)
        results.append(row)
    return results


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #
def write_output(results, out_path, top_n, settings, log=print):
    header = ["zOTU", "Length (bp)", "Best identity (%)", "Status",
              "Assigned species", "Assigned genus", "Species tied at best %"]
    for i in range(1, top_n + 1):
        header += [f"Species {i}", f"Identity {i} (%)", f"Accession {i}"]
    summary = []
    for r in results:
        line = [r["zOTU"], r["length"], r["best_id"], r["status"],
                r["assigned_species"], r["assigned_genus"], r["n_tied"]]
        for sp, (ident, alnlen, qcov, acc) in r["hits"]:
            line += [sp, ident, acc]
        line += [""] * (len(header) - len(line))
        summary.append(line)

    long_header = ["zOTU", "Rank", "Species", "Identity (%)", "Alignment length",
                   "Query coverage (%)", "Accession"]
    long_rows = [[r["zOTU"], i, sp, ident, alnlen, qcov, acc]
                 for r in results
                 for i, (sp, (ident, alnlen, qcov, acc)) in enumerate(r["hits"], 1)]

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        base = os.path.splitext(out_path)[0]
        for suffix, h, rows in (("summary", header, summary), ("top_hits", long_header, long_rows)):
            with open(f"{base}_{suffix}.csv", "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(h)
                w.writerows(rows)
        log(f"openpyxl not installed - wrote {base}_summary.csv and {base}_top_hits.csv")
        return base + "_summary.csv"

    wb = openpyxl.Workbook()
    fills = {"Species": "C6EFCE", "Ambiguous": "FFEB9C", "Genus": "DDEBF7"}
    for title, h, rows in (("Summary", header, summary), ("Top hits", long_header, long_rows)):
        ws = wb.active if title == "Summary" else wb.create_sheet(title)
        ws.title = title
        ws.append(h)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append(row)
            if title == "Summary":
                for key, color in fills.items():
                    if str(row[3]).startswith(key):
                        ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=color)
        ws.freeze_panes = "B2"
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 45)

    ws = wb.create_sheet("Settings")
    for k, v in settings.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    wb.save(out_path)
    log(f"Wrote {out_path}")
    return out_path


# --------------------------------------------------------------------------- #
# Pipeline (used by both the command line and the GUI)
# --------------------------------------------------------------------------- #
DEFAULTS = dict(top=5, threads=os.cpu_count() or 2, species_id=SPECIES_ID, genus_id=GENUS_ID,
                min_id=80.0, min_cov=0.9, tie_delta=0.0, maxaccepts=200, maxrejects=2000,
                trim="auto", db_fasta=None, update_db=False, keep_intermediate=False)


def run_pipeline(fasta, output=None, log=print, progress=None, **options):
    o = dict(DEFAULTS, **options)
    out = output or os.path.splitext(fasta)[0] + "_taxonomy.xlsx"
    base = os.path.splitext(out)[0]
    vsearch = ensure_vsearch(log)

    if o["db_fasta"]:
        db_fasta = o["db_fasta"]
        names = base + "_custom_names.tsv"
        with open(db_fasta) as fi, open(names, "w", newline="") as fn:
            w = csv.writer(fn, delimiter="\t")
            w.writerow(["accession", "species", "description"])
            for line in fi:
                if line.startswith(">"):
                    d = line[1:].strip()
                    w.writerow([d.split()[0], parse_ncbi_name(d), d])
        date = "custom"
    else:
        db_fasta, names = ensure_db(o["update_db"], log)
        date = db_date()

    seqs = read_fasta(fasta)
    if not seqs:
        raise RuntimeError(f"No sequences found in {fasta}")
    order = list(seqs)
    lengths = {k: len(v) for k, v in seqs.items()}
    log(f"{len(order)} zOTUs, length {min(lengths.values())}-{max(lengths.values())} bp")

    if o["trim"] == "auto":
        left, right, notes = detect_trim(seqs)
    elif o["trim"] == "none":
        left, right, notes = 0, 0, []
    else:
        left, right = (int(x) for x in str(o["trim"]).split(","))
        notes = [f"manual: left {left} bp, right {right} bp"]
    query = base + "_query.fasta"
    with open(query, "w") as f:
        for k, s in seqs.items():
            f.write(f">{k}\n{s[left:len(s) - right]}\n")
    log("Primer trimming: " + ("; ".join(notes) if notes else "none"))

    hits = base + "_vsearch_hits.tsv"
    run_vsearch(vsearch, query, db_fasta, hits, o["threads"], o["min_id"], o["min_cov"],
                o["maxaccepts"], o["maxrejects"], log, progress)
    results = classify(hits, names, order, lengths, o["top"], o["species_id"],
                       o["genus_id"], o["tie_delta"])

    counts = defaultdict(int)
    for r in results:
        counts[r["status"].split(":")[0]] += 1
    log("Summary: " + ", ".join(f"{k}: {v}" for k, v in counts.items()))

    settings = {
        "Input": os.path.abspath(fasta),
        "Reference database": db_fasta if o["db_fasta"] else
            "NCBI RefSeq 16S rRNA Targeted Loci (bacteria + archaea)",
        "Database downloaded": date,
        "Aligner": "vsearch --usearch_global (global alignment, terminal gaps ignored), strand both",
        "Species threshold (%)": o["species_id"],
        "Genus threshold (%)": o["genus_id"],
        "Min identity / min query coverage": f"{o['min_id']}% / {o['min_cov']}",
        "Tie tolerance (%)": o["tie_delta"],
        "Primer trimming": "; ".join(notes) if notes else "none",
        "maxaccepts / maxrejects": f"{o['maxaccepts']} / {o['maxrejects']}",
        "Run date": datetime.datetime.now().isoformat(timespec="seconds"),
        "Caveat": "Thresholds were defined for full-length 16S; for short amplicons "
                  "species calls are indicative. Check 'Ambiguous' rows and the Top hits sheet.",
    }
    written = write_output(results, out, o["top"], settings, log)
    if not o["keep_intermediate"]:
        for p in (query, hits) + ((names,) if o["db_fasta"] else ()):
            try:
                os.remove(p)
            except OSError:
                pass
    return written, dict(counts)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("fasta", help="zOTU FASTA file (e.g. zotus.fa from UNOISE3)")
    p.add_argument("-o", "--output", help="output .xlsx (default: <fasta>_taxonomy.xlsx)")
    p.add_argument("--top", type=int, default=DEFAULTS["top"], help="species to report per zOTU (default 5)")
    p.add_argument("--threads", type=int, default=DEFAULTS["threads"])
    p.add_argument("--species-id", type=float, default=SPECIES_ID)
    p.add_argument("--genus-id", type=float, default=GENUS_ID)
    p.add_argument("--min-id", type=float, default=DEFAULTS["min_id"], help="minimum %% identity to keep a hit")
    p.add_argument("--min-cov", "--query-cov", "--query_cov", dest="min_cov", type=float,
                   default=DEFAULTS["min_cov"],
                   help="minimum fraction of the zOTU that must align, passed to vsearch "
                        "--query_cov (0-1, default 0.9)")
    p.add_argument("--tie-delta", type=float, default=0.0,
                   help="species within this %% of the best hit count as tied (default 0 = exactly equal)")
    p.add_argument("--maxaccepts", type=int, default=DEFAULTS["maxaccepts"])
    p.add_argument("--maxrejects", type=int, default=DEFAULTS["maxrejects"])
    p.add_argument("--trim", default="auto",
                   help="primer remnants: 'auto' (default), 'none', or 'LEFT,RIGHT' bases e.g. 7,9")
    p.add_argument("--db-fasta", help="use your own reference FASTA instead of NCBI "
                   "(headers: >accession Genus species ...)")
    p.add_argument("--update-db", action="store_true", help="re-download the NCBI database")
    p.add_argument("--keep-intermediate", action="store_true",
                   help="keep the trimmed query FASTA and raw vsearch hits")
    a = p.parse_args()
    if not 0 < a.min_cov <= 1:
        p.error("--min-cov must be a fraction between 0 and 1 (e.g. 0.99), not a percentage")
    opts = {k: v for k, v in vars(a).items() if k not in ("fasta", "output")}
    try:
        run_pipeline(a.fasta, a.output, **opts)
    except RuntimeError as e:
        sys.exit(str(e))


if __name__ == "__main__":
    main()
